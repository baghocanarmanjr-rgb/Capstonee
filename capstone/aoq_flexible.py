"""Flexible offline AOQ/bidding-result document extraction.

The module accepts common office-document formats, detects tables and labels without
assuming one fixed AOQ template, and returns a standard review payload. Human review
is still required before an AOQ is saved or approved.
"""
from __future__ import annotations

import csv
import json
import os
import re
from collections import OrderedDict
from typing import Any, Dict, Iterable, List, Tuple

from openpyxl import load_workbook
from PIL import Image, ImageFilter, ImageOps
import pytesseract

SUPPORTED_EXTENSIONS = {
    ".pdf", ".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff",
    ".docx", ".xlsx", ".xlsm", ".csv", ".txt",
}

FIELD_ALIASES = {
    "supplier": [
        "supplier", "supplier name", "name of supplier", "bidder", "dealer",
        "company", "vendor", "name of bidder", "winning supplier", "winning bidder",
    ],
    "item": [
        "item", "item name", "article", "particulars", "commodity", "product",
        "description of item", "name of item",
    ],
    "description": [
        "description", "item description", "specification", "specifications", "specs",
        "description specification", "technical specification", "brand model",
    ],
    "qty": ["qty", "quantity", "quantity offered", "quantity quoted", "number of units"],
    "unit": ["unit", "uom", "unit of measure", "measure"],
    "unit_cost": [
        "unit cost", "unit price", "price per unit", "bid price", "quoted price",
        "price", "quotation price",
    ],
    "amount": [
        "amount", "total", "total amount", "grand total", "extended amount",
        "total bid", "bid amount", "quotation amount", "contract amount",
    ],
    "remarks": [
        "remarks", "remark", "evaluation", "recommendation", "result", "status",
        "award basis", "compliance", "finding",
    ],
    "rank": ["rank", "ranking", "lowest", "winner", "award", "awarded"],
}

WINNER_WORDS = (
    "winner", "winning", "awarded", "awardee", "recommended", "selected",
    "lowest calculated responsive", "lowest responsive", "lcrb", "successful bidder",
)

PRICE_WORDS = (
    "unit cost", "unit price", "price per unit", "bid price", "quoted price",
    "price", "amount", "total", "grand total", "bid amount", "quotation amount",
)


def _norm(value: Any) -> str:
    text = str(value if value is not None else "")
    text = text.replace("\u00a0", " ")
    text = re.sub(r"[_/\\|]+", " ", text)
    text = re.sub(r"[^a-zA-Z0-9%₱.+\- ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip().lower()


def _clean(value: Any) -> str:
    text = str(value if value is not None else "")
    text = text.replace("\u00a0", " ")
    return re.sub(r"\s+", " ", text).strip(" \t\r\n:;|-–—")


def _number(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("₱", "").replace("PHP", "").replace(",", "")
    matches = re.findall(r"-?\d+(?:\.\d+)?", text)
    if not matches:
        return 0.0
    try:
        return float(matches[-1])
    except ValueError:
        return 0.0


def _meaningful_row(row: Iterable[Any]) -> bool:
    return any(_clean(cell) for cell in row)


def _trim_rows(rows: List[List[Any]]) -> List[List[Any]]:
    rows = [list(r) for r in rows if _meaningful_row(r)]
    if not rows:
        return []
    max_col = 0
    for row in rows:
        for idx, cell in enumerate(row):
            if _clean(cell):
                max_col = max(max_col, idx + 1)
    return [row[:max_col] + [""] * max(0, max_col - len(row)) for row in rows]


def _table_text(rows: List[List[Any]]) -> str:
    lines = []
    for row in rows:
        values = [_clean(cell) for cell in row]
        if any(values):
            lines.append(" | ".join(values))
    return "\n".join(lines)


def _map_header(value: Any) -> str:
    normalized = _norm(value)
    if not normalized:
        return ""
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            alias_n = _norm(alias)
            if normalized == alias_n or normalized.startswith(alias_n + " ") or normalized.endswith(" " + alias_n):
                return field
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            alias_n = _norm(alias)
            if len(alias_n) >= 4 and alias_n in normalized:
                return field
    return ""


def _configure_tesseract() -> str | None:
    candidates = [
        os.environ.get("TESSERACT_CMD"),
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]
    for candidate in candidates:
        if candidate and os.path.exists(candidate):
            pytesseract.pytesseract.tesseract_cmd = candidate
            return candidate
    return None


def _preprocess_image(image: Image.Image) -> Image.Image:
    img = ImageOps.exif_transpose(image).convert("L")
    if img.width < 2000:
        scale = 2000 / max(img.width, 1)
        img = img.resize((int(img.width * scale), int(img.height * scale)))
    img = ImageOps.autocontrast(img)
    img = img.filter(ImageFilter.SHARPEN)
    return img


def _ocr_image(image: Image.Image) -> str:
    _configure_tesseract()
    processed = _preprocess_image(image)
    variants = []
    for psm in (6, 11):
        try:
            variants.append(pytesseract.image_to_string(processed, config=f"--psm {psm}"))
        except Exception:
            continue
    return max(variants, key=lambda x: len(x.strip()), default="").strip()


def _read_pdf(filepath: str) -> Dict[str, Any]:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("PDF extraction requires PyMuPDF.") from exc

    texts: List[str] = []
    tables: List[Dict[str, Any]] = []
    methods = set()
    warnings: List[str] = []
    with fitz.open(filepath) as doc:
        for page_index, page in enumerate(doc):
            if page_index >= 30:
                warnings.append("Only the first 30 PDF pages were processed.")
                break
            native = (page.get_text("text") or "").strip()
            if len(native) >= 20:
                texts.append(native)
                methods.add("native PDF text")
            else:
                pix = page.get_pixmap(matrix=fitz.Matrix(2.5, 2.5), alpha=False)
                mode = "RGB" if pix.n < 4 else "RGBA"
                image = Image.frombytes(mode, (pix.width, pix.height), pix.samples)
                ocr = _ocr_image(image)
                if ocr:
                    texts.append(ocr)
                    methods.add("OCR on scanned PDF")
            try:
                finder = page.find_tables()
                for table_index, table in enumerate(getattr(finder, "tables", []) or [], start=1):
                    extracted = _trim_rows(table.extract() or [])
                    if extracted:
                        tables.append({
                            "source": f"PDF page {page_index + 1}, table {table_index}",
                            "rows": extracted,
                        })
                        methods.add("PDF table detection")
            except Exception:
                # Table detection is opportunistic; text/OCR remains available.
                pass
    return {
        "raw_text": "\n\n".join(texts).strip(),
        "tables": tables,
        "source_type": "PDF",
        "extraction_method": ", ".join(sorted(methods)) or "PDF extraction",
        "warnings": warnings,
    }


def _read_image(filepath: str) -> Dict[str, Any]:
    with Image.open(filepath) as image:
        text = _ocr_image(image)
    return {
        "raw_text": text,
        "tables": [],
        "source_type": "Scanned image",
        "extraction_method": "Tesseract OCR with image preprocessing",
        "warnings": [],
    }


def _read_docx(filepath: str) -> Dict[str, Any]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Word document extraction requires python-docx.") from exc
    doc = Document(filepath)
    paragraphs = [_clean(p.text) for p in doc.paragraphs if _clean(p.text)]
    tables = []
    for index, table in enumerate(doc.tables, start=1):
        rows = _trim_rows([[cell.text for cell in row.cells] for row in table.rows])
        if rows:
            tables.append({"source": f"Word table {index}", "rows": rows})
    table_text = "\n\n".join(_table_text(t["rows"]) for t in tables)
    return {
        "raw_text": "\n".join(paragraphs + ([table_text] if table_text else [])).strip(),
        "tables": tables,
        "source_type": "Word document",
        "extraction_method": "DOCX paragraphs and tables",
        "warnings": [],
    }


def _read_xlsx(filepath: str) -> Dict[str, Any]:
    wb = load_workbook(filepath, data_only=True, read_only=False)
    tables = []
    text_parts = []
    warnings = []
    for ws in wb.worksheets:
        # Replicate merged-cell values across their range so multi-level supplier headers remain detectable.
        merged_map: Dict[Tuple[int, int], Any] = {}
        for merged in ws.merged_cells.ranges:
            top_value = ws.cell(merged.min_row, merged.min_col).value
            for row in range(merged.min_row, merged.max_row + 1):
                for col in range(merged.min_col, merged.max_col + 1):
                    merged_map[(row, col)] = top_value
        rows = []
        max_row = min(ws.max_row or 0, 10000)
        max_col = min(ws.max_column or 0, 200)
        if ws.max_row and ws.max_row > max_row:
            warnings.append(f"{ws.title}: only the first {max_row} rows were processed.")
        for row_idx in range(1, max_row + 1):
            row = []
            for col_idx in range(1, max_col + 1):
                value = ws.cell(row_idx, col_idx).value
                if value is None and (row_idx, col_idx) in merged_map:
                    value = merged_map[(row_idx, col_idx)]
                row.append(value)
            rows.append(row)
        rows = _trim_rows(rows)
        if rows:
            tables.append({"source": f"Excel sheet: {ws.title}", "rows": rows})
            text_parts.append(f"[{ws.title}]\n{_table_text(rows)}")
    return {
        "raw_text": "\n\n".join(text_parts).strip(),
        "tables": tables,
        "source_type": "Excel workbook",
        "extraction_method": "Dynamic worksheet and cell extraction",
        "warnings": warnings,
    }


def _read_csv(filepath: str) -> Dict[str, Any]:
    rows = []
    with open(filepath, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        sample = fh.read(4096)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(fh, dialect)
        for idx, row in enumerate(reader):
            if idx >= 10000:
                break
            rows.append(row)
    rows = _trim_rows(rows)
    return {
        "raw_text": _table_text(rows),
        "tables": [{"source": "CSV table", "rows": rows}] if rows else [],
        "source_type": "CSV file",
        "extraction_method": "Delimited table extraction",
        "warnings": [],
    }


def _read_txt(filepath: str) -> Dict[str, Any]:
    with open(filepath, "r", encoding="utf-8-sig", errors="replace") as fh:
        text = fh.read()
    return {
        "raw_text": text.strip(),
        "tables": [],
        "source_type": "Text file",
        "extraction_method": "Plain-text extraction",
        "warnings": [],
    }


def extract_document(filepath: str) -> Dict[str, Any]:
    ext = os.path.splitext(filepath)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise RuntimeError("Unsupported file format.")
    if ext == ".pdf":
        result = _read_pdf(filepath)
    elif ext in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}:
        result = _read_image(filepath)
    elif ext == ".docx":
        result = _read_docx(filepath)
    elif ext in {".xlsx", ".xlsm"}:
        result = _read_xlsx(filepath)
    elif ext == ".csv":
        result = _read_csv(filepath)
    else:
        result = _read_txt(filepath)
    result["extension"] = ext
    result["table_count"] = len(result.get("tables") or [])
    if not (result.get("raw_text") or "").strip() and not result.get("tables"):
        raise RuntimeError("No readable text or table was detected in the uploaded document.")
    return result


def _best_header(rows: List[List[Any]]) -> Tuple[int, Dict[int, str], int]:
    best = (-1, {}, 0)
    for idx, row in enumerate(rows[:15]):
        mapping: Dict[int, str] = {}
        for col, cell in enumerate(row):
            field = _map_header(cell)
            if field and field not in mapping.values():
                mapping[col] = field
        unique = set(mapping.values())
        score = len(unique)
        if "supplier" in unique:
            score += 2
        if "item" in unique or "description" in unique:
            score += 2
        if "unit_cost" in unique or "amount" in unique:
            score += 2
        if score > best[2]:
            best = (idx, mapping, score)
    return best


def _record_from_row(row: List[Any], mapping: Dict[int, str]) -> Dict[str, Any]:
    rec: Dict[str, Any] = {}
    for col, field in mapping.items():
        value = row[col] if col < len(row) else ""
        if field in {"qty", "unit_cost", "amount"}:
            rec[field] = _number(value)
        else:
            rec[field] = _clean(value)
    if not rec.get("item") and rec.get("description"):
        rec["item"] = rec["description"][:120]
    if not rec.get("description") and rec.get("item"):
        rec["description"] = rec["item"]
    if not rec.get("amount") and rec.get("qty") and rec.get("unit_cost"):
        rec["amount"] = rec["qty"] * rec["unit_cost"]
    combined = " ".join(str(v) for v in rec.values()).lower()
    rec["is_winner"] = any(word in combined for word in WINNER_WORDS)
    return rec


def _parse_row_table(rows: List[List[Any]], source: str) -> List[Dict[str, Any]]:
    header_idx, mapping, score = _best_header(rows)
    if header_idx < 0 or score < 5:
        return []
    records = []
    last_supplier = ""
    for row in rows[header_idx + 1:]:
        rec = _record_from_row(row, mapping)
        if not rec.get("supplier") and last_supplier and (rec.get("item") or rec.get("description")):
            rec["supplier"] = last_supplier
        if rec.get("supplier"):
            last_supplier = rec["supplier"]
        meaningful = rec.get("supplier") or rec.get("item") or rec.get("description") or rec.get("amount")
        if not meaningful:
            continue
        normalized_join = _norm(" ".join(_clean(x) for x in row))
        if normalized_join in {"total", "grand total"}:
            continue
        rec["source"] = source
        records.append(rec)
    return records


def _metric_from_header(value: str) -> Tuple[str, str]:
    normalized = _norm(value)
    if not normalized:
        return "", ""
    metric = ""
    for candidate in ("unit_cost", "amount", "description", "remarks"):
        aliases = FIELD_ALIASES[candidate]
        for alias in sorted(aliases, key=len, reverse=True):
            alias_n = _norm(alias)
            if alias_n in normalized:
                metric = candidate
                normalized = re.sub(rf"\b{re.escape(alias_n)}\b", " ", normalized)
                break
        if metric:
            break
    if not metric:
        return "", ""
    supplier = re.sub(r"\b(?:supplier|bidder|vendor|quotation|quote|offered|offer)\b", " ", normalized)
    supplier = re.sub(r"\s+", " ", supplier).strip(" -")
    return supplier, metric


def _parse_column_table(rows: List[List[Any]], source: str) -> List[Dict[str, Any]]:
    if len(rows) < 3:
        return []
    max_cols = max(len(r) for r in rows)
    for start in range(min(8, len(rows) - 1)):
        for depth in (2, 3, 1):
            if start + depth >= len(rows):
                continue
            combined = []
            for col in range(max_cols):
                parts = []
                carry = ""
                for ridx in range(start, start + depth):
                    value = _clean(rows[ridx][col] if col < len(rows[ridx]) else "")
                    if value:
                        carry = value
                    elif ridx == start and col > 0:
                        # Common merged-header representation in extracted PDF tables.
                        prev = _clean(rows[ridx][col - 1] if col - 1 < len(rows[ridx]) else "")
                        if prev:
                            carry = prev
                    if carry and (not parts or parts[-1] != carry):
                        parts.append(carry)
                combined.append(" ".join(parts))

            base: Dict[int, str] = {}
            suppliers: Dict[str, Dict[str, int]] = OrderedDict()
            for col, header in enumerate(combined):
                field = _map_header(header)
                if field in {"item", "description", "qty", "unit", "remarks"} and field not in base.values():
                    base[col] = field
                    continue
                supplier, metric = _metric_from_header(header)
                if supplier and metric in {"unit_cost", "amount", "description", "remarks"}:
                    suppliers.setdefault(supplier.title(), {})[metric] = col
            if not suppliers or not ({"item", "description"} & set(base.values())):
                continue

            records = []
            for row in rows[start + depth:]:
                common: Dict[str, Any] = {}
                for col, field in base.items():
                    value = row[col] if col < len(row) else ""
                    common[field] = _number(value) if field == "qty" else _clean(value)
                if not common.get("item") and common.get("description"):
                    common["item"] = common["description"][:120]
                if not common.get("description") and common.get("item"):
                    common["description"] = common["item"]
                if not common.get("item") and not common.get("description"):
                    continue
                for supplier_name, metrics in suppliers.items():
                    rec = dict(common)
                    rec["supplier"] = supplier_name
                    rec["unit_cost"] = _number(row[metrics["unit_cost"]]) if "unit_cost" in metrics and metrics["unit_cost"] < len(row) else 0.0
                    rec["amount"] = _number(row[metrics["amount"]]) if "amount" in metrics and metrics["amount"] < len(row) else 0.0
                    if "description" in metrics and metrics["description"] < len(row):
                        offered = _clean(row[metrics["description"]])
                        if offered:
                            rec["description"] = offered
                    if "remarks" in metrics and metrics["remarks"] < len(row):
                        rec["remarks"] = _clean(row[metrics["remarks"]])
                    if not rec["amount"] and rec.get("qty") and rec["unit_cost"]:
                        rec["amount"] = rec["qty"] * rec["unit_cost"]
                    if rec["unit_cost"] or rec["amount"] or rec.get("description"):
                        remark_text = _norm(rec.get("remarks", ""))
                        supplier_text = _norm(supplier_name)
                        has_winner_hint = any(word in remark_text for word in WINNER_WORDS)
                        supplier_tokens = [t for t in supplier_text.split() if len(t) > 2]
                        supplier_named = bool(supplier_tokens and any(t in remark_text for t in supplier_tokens))
                        # A shared remarks column should not mark every supplier as the winner.
                        rec["is_winner"] = bool(has_winner_hint and supplier_named)
                        rec["source"] = source
                        records.append(rec)
            if records:
                return records
    return []


def _find_labeled_value(lines: List[str], labels: List[str]) -> str:
    label_pattern = "|".join(re.escape(x) for x in labels)
    pattern = re.compile(rf"(?i)^\s*(?:{label_pattern})\s*(?:[:\-–—]|is)?\s*(.+?)\s*$")
    for idx, line in enumerate(lines):
        match = pattern.match(line)
        if match and _clean(match.group(1)):
            return _clean(match.group(1))
        if re.fullmatch(rf"(?i)\s*(?:{label_pattern})\s*[:\-–—]?\s*", line) and idx + 1 < len(lines):
            return _clean(lines[idx + 1])
    return ""


def _text_fallback(raw_text: str, pr_items: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    lines = [_clean(x) for x in str(raw_text or "").splitlines() if _clean(x)]
    joined = "\n".join(lines)
    supplier = _find_labeled_value(lines, FIELD_ALIASES["supplier"])
    item = _find_labeled_value(lines, FIELD_ALIASES["item"])
    description = _find_labeled_value(lines, FIELD_ALIASES["description"])
    qty_text = _find_labeled_value(lines, FIELD_ALIASES["qty"])
    unit = _find_labeled_value(lines, FIELD_ALIASES["unit"])
    unit_cost_text = _find_labeled_value(lines, FIELD_ALIASES["unit_cost"])
    amount_text = _find_labeled_value(lines, FIELD_ALIASES["amount"])
    remarks = _find_labeled_value(lines, FIELD_ALIASES["remarks"])

    if not supplier:
        match = re.search(r"(?i)(?:awarded\s+to|winning\s+(?:supplier|bidder)\s*(?:is|:|-)?)[\s:,-]*([^\n]{2,100})", joined)
        if match:
            supplier = _clean(match.group(1))

    first = pr_items[0] if pr_items else {}
    if not item:
        item = _clean(first.get("Item Name") or first.get("Item Description") or first.get("item") or "")
    if not description:
        description = _clean(first.get("Description") or first.get("description") or item)
    qty = _number(qty_text) or _number(first.get("Quantity") or first.get("qty")) or 1.0
    unit = unit or _clean(first.get("Unit") or first.get("unit") or "unit") or "unit"
    unit_cost = _number(unit_cost_text)
    amount = _number(amount_text)
    if not amount:
        money_values = [_number(x) for x in re.findall(r"(?:₱|PHP|P)\s*[0-9][0-9,]*(?:\.\d{1,2})?", joined, flags=re.I)]
        amount = max(money_values, default=0.0)
    if not amount and qty and unit_cost:
        amount = qty * unit_cost
    if not unit_cost and amount and qty:
        unit_cost = amount / qty
    if not (supplier or item or description or amount):
        return []
    combined = f"{supplier} {remarks} {joined}".lower()
    return [{
        "supplier": supplier or "Unidentified Supplier",
        "item": item or "Winning Offer Item",
        "description": description or item,
        "qty": qty,
        "unit": unit,
        "unit_cost": round(unit_cost, 2),
        "amount": round(amount, 2),
        "remarks": remarks,
        "is_winner": any(word in combined for word in WINNER_WORDS) or bool(supplier),
        "source": "Labeled text / OCR fallback",
    }]



def _explicit_winner_from_text(raw_text: str) -> str:
    lines = [_clean(x) for x in str(raw_text or '').splitlines() if _clean(x)]
    winner = _find_labeled_value(lines, [
        'winning supplier', 'winning bidder', 'awarded supplier', 'awarded to',
        'recommended supplier', 'selected supplier', 'successful bidder',
    ])
    if winner:
        return winner
    joined = '\n'.join(lines)
    match = re.search(
        r'(?i)(?:awarded\s+to|winning\s+(?:supplier|bidder)|recommended\s+(?:supplier|bidder))'
        r'\s*(?:is|:|-)?\s*([^\n]{2,120})', joined
    )
    return _clean(match.group(1)) if match else ''


def _supplier_name_matches(candidate: str, detected: str) -> bool:
    left = _norm(candidate)
    right = _norm(detected)
    if not left or not right:
        return False
    if left in right or right in left:
        return True
    left_tokens = {t for t in left.split() if len(t) > 2 and t not in {'trading','supplies','supply','company','corporation','inc'}}
    right_tokens = {t for t in right.split() if len(t) > 2 and t not in {'trading','supplies','supply','company','corporation','inc'}}
    return bool(left_tokens and right_tokens and left_tokens.intersection(right_tokens))


def _aggregate(records: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    groups: "OrderedDict[str, Dict[str, Any]]" = OrderedDict()
    for rec in records:
        supplier = _clean(rec.get("supplier")) or "Unidentified Supplier"
        key = _norm(supplier)
        group = groups.setdefault(key, {
            "supplier": supplier,
            "amount": 0.0,
            "remarks": "",
            "is_winner": False,
            "items": [],
        })
        item = {
            "item": _clean(rec.get("item")) or "Quoted Item",
            "description": _clean(rec.get("description")) or _clean(rec.get("item")),
            "qty": _number(rec.get("qty")) or 1.0,
            "unit": _clean(rec.get("unit")) or "unit",
            "unit_cost": round(_number(rec.get("unit_cost")), 2),
            "amount": round(_number(rec.get("amount")), 2),
            "remarks": _clean(rec.get("remarks")),
            "source": _clean(rec.get("source")),
        }
        if not item["amount"] and item["qty"] and item["unit_cost"]:
            item["amount"] = round(item["qty"] * item["unit_cost"], 2)
        group["items"].append(item)
        group["amount"] += item["amount"]
        if item["remarks"]:
            group["remarks"] = "; ".join(x for x in [group["remarks"], item["remarks"]] if x)
        group["is_winner"] = bool(group["is_winner"] or rec.get("is_winner"))
    for group in groups.values():
        group["amount"] = round(group["amount"], 2)
    return list(groups.values())


def _pr_fallback_items(pr_items: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    result = []
    for item in pr_items:
        name = _clean(item.get("Item Name") or item.get("Item Description") or item.get("item") or "Requested Item")
        desc = _clean(item.get("Description") or item.get("description") or item.get("Specification") or name)
        qty = _number(item.get("Quantity") or item.get("qty")) or 1.0
        unit = _clean(item.get("Unit") or item.get("unit") or "unit") or "unit"
        result.append({
            "item": name,
            "description": desc,
            "qty": qty,
            "unit": unit,
            "unit_cost": 0.0,
            "amount": 0.0,
            "remarks": "Copied from the linked PR because the uploaded layout did not expose this field clearly.",
            "source": "Linked PR fallback",
        })
    return result


def build_preview(document: Dict[str, Any], pr_items: List[Dict[str, Any]]) -> Dict[str, Any]:
    records: List[Dict[str, Any]] = []
    parsed_sources = []
    for table in document.get("tables") or []:
        rows = _trim_rows(table.get("rows") or [])
        if not rows:
            continue
        row_records = _parse_row_table(rows, table.get("source") or "Detected table")
        column_records = _parse_column_table(rows, table.get("source") or "Detected table")
        chosen = column_records if len(column_records) > len(row_records) else row_records
        if chosen:
            records.extend(chosen)
            parsed_sources.append(table.get("source") or "Detected table")

    if not records:
        records = _text_fallback(document.get("raw_text") or "", pr_items)

    suppliers = _aggregate(records)
    warnings = list(document.get("warnings") or [])
    explicit_winner_name = _explicit_winner_from_text(document.get("raw_text") or "")
    if explicit_winner_name:
        matched_explicit = False
        for supplier in suppliers:
            if _supplier_name_matches(supplier.get("supplier", ""), explicit_winner_name):
                supplier["is_winner"] = True
                matched_explicit = True
            else:
                supplier["is_winner"] = False
        if not matched_explicit:
            warnings.append(f'Winner label detected as "{explicit_winner_name}", but it did not exactly match a parsed supplier row.')
    if not suppliers:
        suppliers = [{
            "supplier": "Unidentified Supplier",
            "amount": 0.0,
            "remarks": "No supplier structure was confidently detected.",
            "is_winner": False,
            "items": _pr_fallback_items(pr_items),
        }]
        warnings.append("The layout did not expose a reliable supplier table. Review and complete all fields manually.")

    explicit_winners = [idx for idx, s in enumerate(suppliers) if s.get("is_winner")]
    winner_index = explicit_winners[0] if len(explicit_winners) == 1 else (0 if len(suppliers) == 1 else -1)
    if len(explicit_winners) > 1:
        warnings.append("More than one supplier looked like a winner. Select the official winning supplier manually.")
    elif winner_index < 0:
        warnings.append("No unambiguous winning supplier was found. Select the official winner during review.")

    offer_items = suppliers[winner_index]["items"] if winner_index >= 0 else (suppliers[0]["items"] if suppliers else [])
    if not offer_items:
        offer_items = _pr_fallback_items(pr_items)
        warnings.append("Winning offer line items were not clearly detected; linked PR items were shown as editable placeholders.")

    found_fields = 0
    total_fields = 6
    winner = suppliers[winner_index] if winner_index >= 0 else {}
    found_fields += int(bool(winner.get("supplier") and winner.get("supplier") != "Unidentified Supplier"))
    found_fields += int(bool(offer_items))
    found_fields += int(any(_number(i.get("qty")) for i in offer_items))
    found_fields += int(any(_clean(i.get("unit")) for i in offer_items))
    found_fields += int(any(_number(i.get("unit_cost")) or _number(i.get("amount")) for i in offer_items))
    found_fields += int(bool(document.get("raw_text") or document.get("tables")))
    confidence = round(found_fields / total_fields * 100, 1)

    table_snapshots = []
    for table in document.get("tables") or []:
        rows = _trim_rows(table.get("rows") or [])
        if rows:
            table_snapshots.append({
                "source": table.get("source") or "Detected table",
                "rows": [[_clean(cell) for cell in row] for row in rows[:100]],
                "truncated": len(rows) > 100,
            })

    return {
        "document_type": document.get("source_type") or "Unknown document",
        "extraction_method": document.get("extraction_method") or "Automatic extraction",
        "confidence": confidence,
        "warnings": list(OrderedDict.fromkeys(warnings)),
        "suppliers": suppliers,
        "winner_index": winner_index,
        "offer_items": offer_items,
        "remarks": _clean(winner.get("remarks")) if winner else "",
        "raw_text": document.get("raw_text") or "",
        "detected_tables": table_snapshots,
        "table_count": document.get("table_count") or len(table_snapshots),
        "parsed_table_sources": parsed_sources,
    }


def safe_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, default=str)
